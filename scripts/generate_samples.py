"""Generate sample PDF/DOCX contracts for manual and automated testing."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt
import fitz

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"


STRUCTURED_PARAS = [
    ("Title", "MUTUAL NON-DISCLOSURE AGREEMENT"),
    (
        "Normal",
        "This Agreement is entered into by Acme Corporation and Jane Doe.",
    ),
    ("Heading 1", "1. Definitions"),
    (
        "Normal",
        '1.1 "Confidential Information" means all non-public information disclosed by either party.',
    ),
    (
        "Normal",
        '1.2 "Receiving Party" means the party receiving Confidential Information.',
    ),
    ("Heading 1", "2. Obligations"),
    (
        "Normal",
        "2.1 The Receiving Party shall protect Confidential Information using reasonable care.",
    ),
    (
        "Normal",
        "2.2 The Receiving Party shall not disclose Confidential Information to third parties without prior written consent.",
    ),
    ("Heading 1", "3. Term and Termination"),
    (
        "Normal",
        "3.1 This Agreement shall remain in effect for a period of three (3) years.",
    ),
    (
        "Normal",
        "3.2 Either party may terminate this Agreement upon thirty (30) days written notice.",
    ),
    ("Heading 1", "4. Governing Law"),
    (
        "Normal",
        "4.1 This Agreement shall be governed by the laws of the State of Delaware.",
    ),
    (
        "Normal",
        "4.2 Notices to Jane Doe at Acme Corporation shall be effective as of January 15, 2024.",
    ),
]

HEADINGLESS = [
    "This is a plain agreement without clear section headings or numbering styles.",
    "The parties agree to keep information confidential for two years from the date of disclosure.",
    "Liability is limited to fees paid in the twelve months preceding the claim.",
    "Notices may be delivered by email with read receipt.",
    "The agreement constitutes the entire understanding between the parties.",
]


def write_docx(path: Path, paragraphs: list[tuple[str, str]] | list[str], structured: bool) -> None:
    doc = Document()
    if structured:
        for style, text in paragraphs:  # type: ignore[misc]
            if style == "Title":
                p = doc.add_paragraph(text)
                p.style = "Title"
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(18)
            elif style.startswith("Heading"):
                p = doc.add_heading(text, level=1)
            else:
                doc.add_paragraph(text)
    else:
        for text in paragraphs:  # type: ignore[assignment]
            doc.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def write_pdf_structured(path: Path) -> None:
    """Create a PDF with larger heading fonts and numbered sections."""
    doc = fitz.open()
    page = doc.new_page()
    y = 72
    lines = [
        (18, True, "MUTUAL NON-DISCLOSURE AGREEMENT"),
        (11, False, "This Agreement is entered into by Acme Corporation and Jane Doe."),
        (14, True, "1. Definitions"),
        (11, False, '1.1 "Confidential Information" means all non-public information disclosed by either party.'),
        (11, False, '1.2 "Receiving Party" means the party receiving Confidential Information.'),
        (14, True, "2. Obligations"),
        (11, False, "2.1 The Receiving Party shall protect Confidential Information using reasonable care."),
        (11, False, "2.2 The Receiving Party shall not disclose Confidential Information to third parties."),
        (14, True, "3. Term and Termination"),
        (11, False, "3.1 This Agreement shall remain in effect for a period of three (3) years."),
        (11, False, "3.2 Either party may terminate this Agreement upon thirty (30) days written notice."),
        (14, True, "4. Governing Law"),
        (11, False, "4.1 This Agreement shall be governed by the laws of the State of Delaware."),
        (11, False, "4.2 Notices to Jane Doe at Acme Corporation shall be effective as of January 15, 2024."),
    ]
    for size, bold, text in lines:
        font = "helv"
        # PyMuPDF insert_text uses built-in fonts; bold via fontname
        fname = "hebo" if bold else "helv"
        page.insert_text((72, y), text, fontsize=size, fontname=fname)
        y += size + 10
        if y > 750:
            page = doc.new_page()
            y = 72
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()


def write_pdf_headingless(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for text in HEADINGLESS:
        page.insert_text((72, y), text, fontsize=11, fontname="helv")
        y += 28
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()


def main() -> None:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    write_docx(SAMPLES / "structured_nda.docx", STRUCTURED_PARAS, structured=True)
    write_docx(SAMPLES / "headingless_policy.docx", HEADINGLESS, structured=False)
    write_pdf_structured(SAMPLES / "structured_nda.pdf")
    write_pdf_headingless(SAMPLES / "headingless_policy.pdf")

    # Second version of NDA for future drift tests (minor wording change)
    v2 = list(STRUCTURED_PARAS)
    v2 = [
        (
            style,
            text.replace("three (3) years", "five (5) years").replace(
                "shall protect", "shall use commercially reasonable efforts to protect"
            ),
        )
        for style, text in STRUCTURED_PARAS
    ]
    write_docx(SAMPLES / "structured_nda_v2.docx", v2, structured=True)
    print(f"Wrote samples to {SAMPLES}")


if __name__ == "__main__":
    main()
